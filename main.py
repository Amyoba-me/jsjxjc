import os
import io
import asyncio
import random
import re
import logging
from html.parser import HTMLParser

import aiohttp
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from aiogram import Bot, Dispatcher, types
from aiogram import F
from aiogram.filters import Command
from aiogram.types import BufferedInputFile, ReplyKeyboardMarkup, KeyboardButton, ReplyKeyboardRemove
from aiogram.fsm.storage.memory import MemoryStorage

BOT_TOKEN = os.getenv("BOT_TOKEN")

if not BOT_TOKEN:
    raise ValueError(
        "Переменная окружения BOT_TOKEN не найдена! "
        "Убедитесь, что вы указали BOT_TOKEN в настройках хостинга."
    )

logging.basicConfig(level=logging.INFO)

# 1. Цветовая гамма племен и запахов
SMELL_MAP = {
    "#dfdc8f": "Племя Ветра 🌾",
    "#ff861c": "Племя Солнца ☀️",
    "#00b4d8": "Племя Потока 🌊",
    "#71c68b": "Племя Мрака 🌲",
    "#576198": "Клан Горных Вершин 🏔️",
    "#befffb": "Звёздные Угодья ✨",
    "#f777a6": "Домашки 🏠",
    "#e3d1c8": "Одиночки 🐾",
    "#911922": "Сумеречный Лес 🌑",
}

active_tasks = {}

def parse_color_to_name(color_raw: str) -> str:
    """Определяет запах племени по HEX или RGB строке."""
    if not color_raw:
        return "Неизвестно"

    color_raw = color_raw.strip().lower()

    rgb_match = re.search(r"rgb\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)\s*\)", color_raw)
    if rgb_match:
        r, g, b = map(int, rgb_match.groups())
        hex_code = f"#{r:02x}{g:02x}{b:02x}"
        return SMELL_MAP.get(hex_code, f"Неизвестный запах ({hex_code})")

    hex_match = re.search(r"#[0-9a-f]{6}\b", color_raw)
    if hex_match:
        hex_code = hex_match.group(0)
        return SMELL_MAP.get(hex_code, f"Неизвестный запах ({hex_code})")

    if color_raw in SMELL_MAP:
        return SMELL_MAP[color_raw]

    return color_raw

class SimpleHTMLTextExtractor(HTMLParser):
    def __init__(self):
        super().__init__()
        self.result = []

    def handle_data(self, data):
        cleaned = data.strip()
        if cleaned:
            self.result.append(cleaned)

def extract_smell_from_html(html_content: str) -> str:
    """Извлечение запаха из style-тегов."""
    smell_pattern = r"Запах племени:\s*<span[^>]*style=[\"']([^\"']+)[\"'][^>]*>"
    match = re.search(smell_pattern, html_content, re.IGNORECASE)
    if match:
        style_str = match.group(1)
        color_match = re.search(r"background(?:-color)?:\s*([^;]+)", style_str, re.I) or \
                      re.search(r"color:\s*([^;]+)", style_str, re.I)
        if color_match:
            return parse_color_to_name(color_match.group(1).strip())
    return "Неизвестно"

def parse_character_data(html_content: str, char_id: int) -> dict:
    """Парсинг имени, запаха, родителей, боевых умений и часов в игре."""
    parser = SimpleHTMLTextExtractor()
    parser.feed(html_content)
    tokens = parser.result

    name = f"Персонаж {char_id}"
    if "Общий рейтинг:" in tokens:
        idx = tokens.index("Общий рейтинг:")
        if idx >= 2:
            name = tokens[idx - 2]

    smell = extract_smell_from_html(html_content)

    # Парсинг родителей
    parents = []
    link_matches = re.finditer(r'<a[^>]+href=["\'](?:https?://[^/]+)?/p/(\d+)["\'][^>]*>(.*?)</a>', html_content, re.IGNORECASE | re.DOTALL)
    
    for match in link_matches:
        target_id = int(match.group(1))
        raw_text = match.group(2)
        target_name = re.sub(r'<[^>]+>', '', raw_text).strip() or f"ID {target_id}"

        start_pos = max(0, match.start() - 150)
        end_pos = min(len(html_content), match.end() + 150)
        snippet = html_content[start_pos:end_pos].lower()

        if any(w in snippet for w in ["родител", "мать", "отец", "родители"]):
            parents.append({"id": target_id, "name": target_name})

    unique_parents = {}
    for p in parents:
        unique_parents[p["id"]] = p["name"]

    # Парсинг боевых умений и времени в игре
    combat_skills = 0
    combat_match = re.search(r"Боевые\s+умения:\s*(\d+)", html_content, re.IGNORECASE)
    if combat_match:
        combat_skills = int(combat_match.group(1))

    game_hours = 0.0
    time_match = re.search(r"(?:Время в игре|Игровое время):\s*([\d\.,]+)\s*(?:ч|часов|ч\.)", html_content, re.IGNORECASE)
    if time_match:
        try:
            game_hours = float(time_match.group(1).replace(",", "."))
        except ValueError:
            game_hours = 0.0

    return {
        "id": char_id,
        "name": name,
        "smell": smell,
        "parents": unique_parents,
        "combat_skills": combat_skills,
        "game_hours": game_hours,
    }

async def fetch_page(session: aiohttp.ClientSession, char_id: int, max_retries: int = 2) -> str | None:
    """Загрузка страницы с повторными попытками."""
    url = f"https://stats.worldcats.ru/p/{char_id}"
    timeout = aiohttp.ClientTimeout(total=10.0)

    for attempt in range(1, max_retries + 1):
        try:
            async with session.get(url, timeout=timeout) as response:
                if response.status == 200:
                    return await response.text()
                elif response.status in (404, 403):
                    return None
        except (aiohttp.ClientError, asyncio.TimeoutError):
            pass
        
        if attempt < max_retries:
            await asyncio.sleep(1.0)

    return None

def get_reply_keyboard(is_paused: bool = False) -> ReplyKeyboardMarkup:
    """Клавиатура управления парсингом."""
    pause_btn_text = "▶️ Продолжить" if is_paused else "⏸ Пауза"
    return ReplyKeyboardMarkup(
        keyboard=[
            [
                KeyboardButton(text=pause_btn_text),
                KeyboardButton(text="📥 Промежуточный результат"),
                KeyboardButton(text="⏹ Остановить")
            ]
        ],
        resize_keyboard=True
    )

def build_children_map(characters_data: dict) -> dict:
    """Карта детей для каждого родителя."""
    children_map = {}
    for char_id, char_info in characters_data.items():
        for parent_id in char_info["parents"].keys():
            if parent_id not in children_map:
                children_map[parent_id] = []
            children_map[parent_id].append({
                "id": char_id,
                "name": char_info["name"]
            })
    return children_map

def generate_txt_report(state: dict, is_interim: bool = False) -> BufferedInputFile:
    """Генерация .txt отчета (итогового или промежуточного)."""
    start_id = state["start_id"]
    end_id = state["end_id"]
    current_id = state["current_id"]
    characters_data = state["characters_data"]
    failed_ids = state["failed_ids"]
    smell_counts = state["smell_counts"]

    report = []
    report.append("══════════════════════════════════════════════════")
    if is_interim:
        report.append(f"   ПРОМЕЖУТОЧНЫЙ ОТЧЕТ [Текущий ID: {current_id} | Диапазон: {start_id} - {end_id}]")
    else:
        report.append(f"   ОТЧЕТ ПАРСИНГА ДИАПАЗОНА [{start_id} - {end_id}]")
    report.append("══════════════════════════════════════════════════\n")
    
    # 1. Количество ошибок
    report.append(f"Всего обработано персонажей: {len(characters_data)}")
    report.append(f"Всего ID с ошибкой/пропущенных: {len(failed_ids)}\n")

    # 2. Топ племен по количеству ID
    report.append("--- ТОП ПЛЕМЕН (ЗАПАХОВ) ПО КОЛИЧЕСТВУ ПЕРСОНАЖЕЙ ---")
    if smell_counts:
        sorted_smells = sorted(smell_counts.items(), key=lambda x: x[1], reverse=True)
        for idx, (smell_name, count) in enumerate(sorted_smells, 1):
            report.append(f"{idx}. {smell_name}: {count} перс.")
    else:
        report.append("Данные отсутствуют.")
    report.append("\n")

    # 3. Топ 100 по боевым умениям
    report.append("--- ТОП-100 ПО БОЕВЫМ УМЕНИЯМ ---")
    sorted_by_combat = sorted(
        characters_data.values(),
        key=lambda x: x.get("combat_skills", 0),
        reverse=True
    )[:100]

    if sorted_by_combat and any(c.get("combat_skills", 0) > 0 for c in sorted_by_combat):
        for idx, c in enumerate(sorted_by_combat, 1):
            report.append(f"{idx}. ID {c['id']} - {c['name']} ({c['smell']}) — {c['combat_skills']} ед.")
    else:
        report.append("Нет данных или у всех персонажей 0 ед.")
    report.append("\n")

    # 4. Топ 100 по часам в игре
    report.append("--- ТОП-100 ПО ВРЕМЕНИ В ИГРЕ ---")
    sorted_by_hours = sorted(
        characters_data.values(),
        key=lambda x: x.get("game_hours", 0.0),
        reverse=True
    )[:100]

    if sorted_by_hours and any(c.get("game_hours", 0.0) > 0 for c in sorted_by_hours):
        for idx, c in enumerate(sorted_by_hours, 1):
            report.append(f"{idx}. ID {c['id']} - {c['name']} ({c['smell']}) — {c['game_hours']} ч.")
    else:
        report.append("Нет данных или у всех персонажей 0 ч.")
    report.append("\n")

    file_bytes = "\n".join(report).encode('utf-8')
    prefix = "interim_report" if is_interim else "report"
    return BufferedInputFile(file_bytes, filename=f"{prefix}_{start_id}_{end_id}.txt")

def generate_excel_report(state: dict, is_interim: bool = False) -> BufferedInputFile:
    """Генерация Excel файла: ID - Имя - Запах племени - Родители - Дети."""
    start_id = state["start_id"]
    end_id = state["end_id"]
    characters_data = state["characters_data"]
    children_map = build_children_map(characters_data)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Персонажи"
    ws.views.sheetView[0].showGridLines = True

    # Стилизация Excel
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=11, color="000000")
    fill_header = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    fill_zebra = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    headers = ["айди", "Имя", "запах племени", "Родители", "Дети"]
    
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
    ws.row_dimensions[1].height = 25

    row_idx = 2
    for char_id in sorted(characters_data.keys()):
        char_info = characters_data[char_id]

        # Строка родителей
        parents = char_info["parents"]
        if parents:
            parents_str = ", ".join([f"{p_name} (ID: {p_id})" for p_id, p_name in parents.items()])
        else:
            parents_str = "—"

        # Строка детей
        children = children_map.get(char_id, [])
        if children:
            children_str = ", ".join([f"{c['name']} (ID: {c['id']})" for c in children])
        else:
            children_str = "—"

        row_values = [
            char_id,
            char_info["name"],
            char_info["smell"],
            parents_str,
            children_str
        ]

        row_fill = fill_zebra if row_idx % 2 == 0 else fill_white
        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_data
            cell.fill = row_fill
            cell.border = thin_border
            if col_idx == 1:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')

        ws.row_dimensions[row_idx].height = 20
        row_idx += 1

    # Подгонка ширины столбцов
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    prefix = "interim_characters" if is_interim else "characters"
    return BufferedInputFile(stream.read(), filename=f"{prefix}_{start_id}_{end_id}.xlsx")

bot = Bot(token=BOT_TOKEN)
dp = Dispatcher(storage=MemoryStorage())

async def run_parser_task(chat_id: int, start_id: int, end_id: int, batch_size: int = 100):
    state = active_tasks[chat_id]
    total_ids = end_id - start_id + 1

    await bot.send_message(
        chat_id,
        f"🚀 **Запуск парсинга диапазонов {start_id} – {end_id}** (всего ID: {total_ids})",
        parse_mode="Markdown",
        reply_markup=get_reply_keyboard()
    )

    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/124.0.0.0 Safari/537.36"
        ),
        "Accept-Language": "ru-RU,ru;q=0.9,en-US;q=0.8,en;q=0.7",
    }

    async with aiohttp.ClientSession(headers=headers) as session:
        current_id = state["current_id"]

        while current_id <= end_id and state["is_running"]:
            if state["jump_to_id"] is not None:
                current_id = state["jump_to_id"]
                state["jump_to_id"] = None
                if current_id > end_id:
                    break

            while state["is_paused"] and state["is_running"]:
                await asyncio.sleep(1.0)
                if not state["is_running"]:
                    break

            if not state["is_running"]:
                break

            current_batch_end = min(current_id + batch_size - 1, end_id)
            batch_failed_ids = []

            for cid in range(current_id, current_batch_end + 1):
                if state["jump_to_id"] is not None or not state["is_running"]:
                    break

                while state["is_paused"] and state["is_running"]:
                    await asyncio.sleep(1.0)

                if not state["is_running"]:
                    break

                state["current_id"] = cid
                html_content = await fetch_page(session, cid, max_retries=2)
                
                if html_content:
                    try:
                        data = parse_character_data(html_content, cid)
                        state["characters_data"][cid] = data
                        
                        smell = data["smell"]
                        state["smell_counts"][smell] = state["smell_counts"].get(smell, 0) + 1
                    except Exception as e:
                        logging.error(f"❌ Ошибка разбора ID {cid}: {e}")
                        state["failed_ids"].append(cid)
                        batch_failed_ids.append(cid)
                else:
                    logging.warning(f"⚠️ ID {cid} не загрузился или не существует")
                    state["failed_ids"].append(cid)
                    batch_failed_ids.append(cid)

                await asyncio.sleep(random.uniform(0.3, 0.6))

            if state["jump_to_id"] is not None:
                continue

            if not state["is_running"]:
                break

            processed_count = len(state["characters_data"]) + len(state["failed_ids"])
            progress_percent = min(100, int((processed_count / total_ids) * 100))

            report_batch_text = (
                f"📊 **Промежуточный отчёт**\n"
                f"Обработан пакет: `id{current_id}` — `id{current_batch_end}`\n"
                f"Прогресс: `{processed_count}/{total_ids}` ({progress_percent}%)\n"
                f"Ошибок в пакете: {len(batch_failed_ids)}"
            )
            
            await bot.send_message(
                chat_id,
                report_batch_text,
                parse_mode="Markdown"
            )

            current_id = current_batch_end + 1

    # Формируем итоговые файлы
    txt_file = generate_txt_report(state, is_interim=False)
    excel_file = generate_excel_report(state, is_interim=False)
    status_caption = "⏹ Парсинг остановлен пользователем!" if not state["is_running"] else "✅ Парсинг завершён!"

    # Отправляем TXT и Excel файлы
    await bot.send_document(
        chat_id,
        txt_file,
        caption=f"{status_caption}\nТекстовый отчёт (диапазон: `id{start_id}` — `id{end_id}`).",
        parse_mode="Markdown",
        reply_markup=ReplyKeyboardRemove()
    )

    await bot.send_document(
        chat_id,
        excel_file,
        caption=f"📊 Таблица Excel с персонажами.",
        parse_mode="Markdown"
    )

    active_tasks.pop(chat_id, None)

@dp.message(Command("start"))
async def cmd_start(message: types.Message):
    welcome_text = (
        "👋 **Привет! Я бот-парсер персонажей worldcats.ru**\n\n"
        "📌 **Управление парсингом:**\n"
        "• `/parse [старт_id] [конец_id]` — запустить парсинг\n"
        "• `/resume [id]` — вернуться/перейти к конкретному ID без потери прогресса\n\n"
        "По завершении бот отправит `.txt` с топами и отчётом, а также файл `.xlsx`."
    )
    await message.answer(welcome_text, parse_mode="Markdown", reply_markup=ReplyKeyboardRemove())

@dp.message(Command("parse"))
async def cmd_parse(message: types.Message):
    chat_id = message.chat.id

    if chat_id in active_tasks and active_tasks[chat_id]["is_running"]:
        await message.answer("⚠️ У вас уже запущен парсинг! Используйте кнопки клавиатуры внизу.")
        return

    args = message.text.split()
    if len(args) < 3:
        await message.answer("Укажите старт и конец диапазона.\nПример: `/parse 1 500`", parse_mode="Markdown")
        return

    try:
        start_id = int(args[1])
        end_id = int(args[2])
    except ValueError:
        await message.answer("❌ ID должны быть целыми числами!")
        return

    if start_id > end_id:
        await message.answer("❌ Стартовый ID не может быть больше конечного!")
        return

    active_tasks[chat_id] = {
        "is_running": True,
        "is_paused": False,
        "start_id": start_id,
        "end_id": end_id,
        "current_id": start_id,
        "characters_data": {},
        "failed_ids": [],
        "smell_counts": {},
        "jump_to_id": None
    }

    asyncio.create_task(run_parser_task(chat_id, start_id, end_id))

@dp.message(Command("resume"))
async def cmd_resume(message: types.Message):
    chat_id = message.chat.id

    if chat_id not in active_tasks or not active_tasks[chat_id]["is_running"]:
        await message.answer("❌ В данный момент у вас нет активного парсинга.")
        return

    args = message.text.split()
    if len(args) < 2:
        await message.answer("Укажите ID, к которому нужно вернуться/перейти.\nПример: `/resume 450`", parse_mode="Markdown")
        return

    try:
        target_id = int(args[1])
    except ValueError:
        await message.answer("❌ ID должен быть целым числом!")
        return

    state = active_tasks[chat_id]
    if target_id < state["start_id"] or target_id > state["end_id"]:
        await message.answer(f"❌ Указанный ID вне рамок текущего диапазона ({state['start_id']} – {state['end_id']})!")
        return

    state["jump_to_id"] = target_id
    state["is_paused"] = False
    await message.answer(f"🔄 Возвращаемся к парсингу с **ID {target_id}**...", parse_mode="Markdown", reply_markup=get_reply_keyboard(False))

@dp.message(F.text.in_(["⏸ Пауза", "▶️ Продолжить"]))
async def process_toggle_pause(message: types.Message):
    chat_id = message.chat.id

    if chat_id not in active_tasks or not active_tasks[chat_id]["is_running"]:
        await message.answer("У вас нет активного парсинга.", reply_markup=ReplyKeyboardRemove())
        return

    state = active_tasks[chat_id]
    state["is_paused"] = not state["is_paused"]

    status_text = "⏸ Парсинг поставлен на паузу" if state["is_paused"] else "▶️ Парсинг возобновлён"
    await message.answer(status_text, reply_markup=get_reply_keyboard(state["is_paused"]))

@dp.message(F.text == "📥 Промежуточный результат")
async def process_send_interim_report(message: types.Message):
    chat_id = message.chat.id

    if chat_id not in active_tasks or not active_tasks[chat_id]["is_running"]:
        await message.answer("У вас нет активного парсинга.", reply_markup=ReplyKeyboardRemove())
        return

    state = active_tasks[chat_id]

    if not state["characters_data"]:
        await message.answer("⚠️ Ещё не собрано ни одной записи для отчёта.")
        return

    await message.answer("⏳ Формирую текущие промежуточные отчёты...")

    # Генерация промежуточных файлов
    txt_file = generate_txt_report(state, is_interim=True)
    excel_file = generate_excel_report(state, is_interim=True)

    current_id = state["current_id"]
    total = state["end_id"] - state["start_id"] + 1

    await bot.send_document(
        chat_id,
        txt_file,
        caption=f"📥 **Промежуточный текстовый отчёт**\nСобрано к ID: `{current_id}` (из {total})",
        parse_mode="Markdown"
    )

    await bot.send_document(
        chat_id,
        excel_file,
        caption=f"📊 **Промежуточная Excel таблица**",
        parse_mode="Markdown"
    )

@dp.message(F.text == "⏹ Остановить")
async def process_stop_parser(message: types.Message):
    chat_id = message.chat.id

    if chat_id not in active_tasks or not active_tasks[chat_id]["is_running"]:
        await message.answer("У вас нет активного парсинга.", reply_markup=ReplyKeyboardRemove())
        return

    state = active_tasks[chat_id]
    state["is_running"] = False
    await message.answer("⏳ Останавливаю парсинг и подготавливаю итоговые файлы...")

async def main():
    await dp.start_polling(bot)

if __name__ == "__main__":
    asyncio.run(main())